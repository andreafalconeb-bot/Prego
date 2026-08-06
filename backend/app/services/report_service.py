import io
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .. import models

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ANOMALY_FILL = PatternFill(start_color="FDECEC", end_color="FDECEC", fill_type="solid")

COLUMNS = [
    ("Placa", 14),
    ("Fecha entrada", 14),
    ("Hora entrada", 12),
    ("Fecha salida", 14),
    ("Hora salida", 12),
    ("Duración", 12),
    ("Cámara entrada", 16),
    ("Cámara salida", 16),
    ("N° Ticket", 16),
    ("Ticket impreso", 14),
    ("Reingresos (misma estadía)", 12),
    ("Estado", 12),
    ("Anomalía", 30),
]


def _duration_str(entry: datetime, exit_: datetime | None, until: datetime) -> str:
    end = exit_ or until
    delta = end - entry
    total_minutes = int(delta.total_seconds() // 60)
    hours, minutes = divmod(total_minutes, 60)
    return f"{hours:02d}:{minutes:02d}"


def _visits_overlapping_range(db: Session, start: datetime, end: datetime) -> list[models.Visit]:
    return (
        db.query(models.Visit)
        .filter(
            models.Visit.entry_time < end,
            (models.Visit.exit_time.is_(None)) | (models.Visit.exit_time >= start),
        )
        .order_by(models.Visit.entry_time)
        .all()
    )


def _write_visit_rows(ws, visits: list[models.Visit], cameras_by_id: dict[int, models.Camera], now: datetime):
    for row_idx, visit in enumerate(visits, start=2):
        entry_camera = cameras_by_id.get(visit.entry_camera_id)
        exit_camera = cameras_by_id.get(visit.exit_camera_id)
        values = [
            visit.plate,
            visit.entry_time.strftime("%Y-%m-%d"),
            visit.entry_time.strftime("%H:%M:%S"),
            visit.exit_time.strftime("%Y-%m-%d") if visit.exit_time else "",
            visit.exit_time.strftime("%H:%M:%S") if visit.exit_time else "",
            _duration_str(visit.entry_time, visit.exit_time, now),
            entry_camera.name if entry_camera else "",
            exit_camera.name if exit_camera else "",
            visit.ticket_number or "",
            "Sí" if visit.ticket_printed else "No",
            visit.reentry_count,
            "Dentro" if visit.status == models.VisitStatus.OPEN else "Salió",
            visit.anomaly_note or "",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if visit.anomaly:
                cell.fill = ANOMALY_FILL


def _new_sheet(wb: Workbook, title: str):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(title)
    ws.title = title
    for col_idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    return ws


def build_daily_report(db: Session, day: date) -> bytes:
    start = datetime.combine(day, time.min)
    end = start + timedelta(days=1)
    visits = _visits_overlapping_range(db, start, end)
    cameras_by_id = {c.id: c for c in db.query(models.Camera).all()}

    wb = Workbook()
    ws = _new_sheet(wb, day.strftime("%Y-%m-%d"))
    _write_visit_rows(ws, visits, cameras_by_id, datetime.utcnow())
    if ws.auto_filter:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_range_report(db: Session, start_day: date, end_day: date) -> bytes:
    cameras_by_id = {c.id: c for c in db.query(models.Camera).all()}
    now = datetime.utcnow()

    wb = Workbook()
    wb.remove(wb.active)
    day = start_day
    while day <= end_day:
        start = datetime.combine(day, time.min)
        end = start + timedelta(days=1)
        visits = _visits_overlapping_range(db, start, end)
        ws = _new_sheet(wb, day.strftime("%Y-%m-%d"))
        _write_visit_rows(ws, visits, cameras_by_id, now)
        if ws.auto_filter:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"
        day += timedelta(days=1)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
