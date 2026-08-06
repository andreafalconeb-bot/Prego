import io
from datetime import date, datetime

from openpyxl import load_workbook

from app.services import visit_service
from app.services.report_service import build_daily_report


def test_daily_report_lists_visits_for_the_day(db, entry_camera, exit_camera):
    visit_service.handle_detection(db, entry_camera, "ABC123", 0.9, datetime(2026, 1, 1, 8, 0), None)
    visit_service.handle_detection(db, exit_camera, "ABC123", 0.9, datetime(2026, 1, 1, 9, 0), None)
    # Different day: must NOT appear in the Jan 1 report.
    visit_service.handle_detection(db, entry_camera, "ZZZ999", 0.9, datetime(2026, 1, 2, 8, 0), None)

    content = build_daily_report(db, date(2026, 1, 1))
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][0] == "ABC123"
    assert rows[0][1] == "2026-01-01"
    assert rows[0][2] == "08:00:00"
    assert rows[0][4] == "09:00:00"


def test_daily_report_header_present(db):
    content = build_daily_report(db, date(2026, 1, 1))
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert "Placa" in header
    assert "N° Ticket" in header
